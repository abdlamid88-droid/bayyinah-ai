#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
بيّنة AI — Bayyinah Engine
Presentation Generator for 'تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي 2026'
Generates a professional, modern 7-slide presentation in 16:9 widescreen.
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# COLOR PALETTE & DESIGN SYSTEM
# ==============================================================================
COLOR_NAVY_DARK      = RGBColor(15, 23, 42)      # #0F172A (Deep Navy/Slate 900)
COLOR_NAVY_CARD      = RGBColor(30, 41, 59)      # #1E293B (Slate 800)
COLOR_NAVY_BORDER    = RGBColor(51, 65, 85)      # #334155 (Slate 700)
COLOR_EMERALD        = RGBColor(5, 150, 105)     # #059669 (Islamic Emerald 600)
COLOR_EMERALD_BRIGHT = RGBColor(16, 185, 129)    # #10B981 (Emerald 500)
COLOR_EMERALD_BG     = RGBColor(236, 253, 245)   # #ECFDF5 (Emerald 50)
COLOR_EMERALD_BORDER = RGBColor(167, 243, 208)   # #A7F3D0 (Emerald 200)
COLOR_BG_LIGHT       = RGBColor(248, 250, 252)   # #F8FAFC (Slate 50)
COLOR_CARD_LIGHT     = RGBColor(255, 255, 255)   # #FFFFFF (Pure White)
COLOR_BORDER_LIGHT   = RGBColor(226, 232, 240)   # #E2E8F0 (Slate 200)
COLOR_TEXT_DARK      = RGBColor(15, 23, 42)      # #0F172A (Slate 900)
COLOR_TEXT_BODY      = RGBColor(51, 65, 85)      # #334155 (Slate 700)
COLOR_TEXT_MUTED     = RGBColor(100, 116, 139)   # #64748B (Slate 500)
COLOR_WHITE          = RGBColor(255, 255, 255)
COLOR_AMBER          = RGBColor(217, 119, 6)     # #D97706 (Amber 600)
COLOR_AMBER_BG       = RGBColor(254, 243, 199)   # #FEF3C7 (Amber 100)
COLOR_BLUE_BG        = RGBColor(239, 246, 255)   # #EFF6FF (Blue 50)
COLOR_BLUE_BORDER    = RGBColor(191, 219, 254)   # #BFDBFE (Blue 200)
COLOR_BLUE_TEXT      = RGBColor(29, 78, 216)     # #1D4ED8 (Blue 700)

FONT_HEADING = "Noto Sans Arabic"
FONT_BODY    = "Noto Sans Arabic"
FONT_LATIN   = "Arial"

# ==============================================================================
# HELPER FUNCTIONS
# ==============================================================================

def set_rtl(paragraph):
    """Enforces Right-to-Left paragraph alignment and OpenXML RTL attribute."""
    paragraph.alignment = PP_ALIGN.RIGHT
    try:
        paragraph._p.get_or_add_pPr().set('rtl', '1')
    except Exception:
        pass

def set_ltr(paragraph):
    """Enforces Left-to-Right paragraph alignment."""
    paragraph.alignment = PP_ALIGN.LEFT
    try:
        paragraph._p.get_or_add_pPr().set('rtl', '0')
    except Exception:
        pass

def set_center(paragraph, is_rtl=True):
    """Enforces Centered paragraph alignment."""
    paragraph.alignment = PP_ALIGN.CENTER
    if is_rtl:
        try:
            paragraph._p.get_or_add_pPr().set('rtl', '1')
        except Exception:
            pass

def create_full_slide(prs, bg_color):
    """Creates a blank slide with a full background color rect."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    bg.line.fill.background()
    return slide

def add_header(slide, tag_text, title_text, subtitle_text, is_dark=False):
    """Adds a standardized top header section with badge, main title, and subtitle."""
    left = Inches(0.8)
    top = Inches(0.46)
    width = Inches(11.733)
    height = Inches(1.4)

    tx_box = slide.shapes.add_textbox(left, top, width, height)
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    # 1. Kicker / Tag
    p_tag = tf.paragraphs[0]
    p_tag.text = tag_text
    set_rtl(p_tag)
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_EMERALD_BRIGHT if is_dark else COLOR_EMERALD
    p_tag.space_after = Pt(2)

    # 2. Main Title
    p_title = tf.add_paragraph()
    p_title.text = title_text
    set_rtl(p_title)
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE if is_dark else COLOR_TEXT_DARK
    p_title.space_after = Pt(3)

    # 3. Subtitle
    p_sub = tf.add_paragraph()
    p_sub.text = subtitle_text
    set_rtl(p_sub)
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = RGBColor(148, 163, 184) if is_dark else COLOR_TEXT_MUTED

def add_footer(slide, current_num, total_num=7, is_dark=False):
    """Adds a clean two-part footer: RTL project text on the right, LTR page number on the left."""
    # Divider line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.015)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(51, 65, 85) if is_dark else COLOR_BORDER_LIGHT
    line.line.fill.background()

    # Right Box: Project & Event text (Arabic RTL)
    tx_right = slide.shapes.add_textbox(Inches(2.5), Inches(6.98), Inches(10.033), Inches(0.35))
    tf_r = tx_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.text = "بيّنة AI — Bayyinah Engine  |  تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي 2026"
    set_rtl(p_r)
    p_r.font.name = FONT_BODY
    p_r.font.size = Pt(10)
    p_r.font.color.rgb = RGBColor(148, 163, 184) if is_dark else COLOR_TEXT_MUTED

    # Left Box: Slide Number (Strict LTR to prevent BiDi inversion)
    tx_left = slide.shapes.add_textbox(Inches(0.8), Inches(6.98), Inches(1.5), Inches(0.35))
    tf_l = tx_left.text_frame
    tf_l.word_wrap = False
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = f"{current_num:02d} / {total_num:02d}"
    set_ltr(p_l)
    p_l.font.name = FONT_LATIN
    p_l.font.size = Pt(10)
    p_l.font.bold = True
    p_l.font.color.rgb = COLOR_EMERALD_BRIGHT if is_dark else COLOR_EMERALD

def build_presentation(output_path="Bayyinah_AI_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    print("Building Slide 1: Cover...")
    build_slide_1_cover(prs)

    print("Building Slide 2: Problem & Context...")
    build_slide_2_problem(prs)

    print("Building Slide 3: Solution & Value Proposition...")
    build_slide_3_solution(prs)

    print("Building Slide 4: Technical Architecture...")
    build_slide_4_architecture(prs)

    print("Building Slide 5: Reliability & Safety...")
    build_slide_5_safety(prs)

    print("Building Slide 6: Execution Roadmap...")
    build_slide_6_roadmap(prs)

    print("Building Slide 7: Team & Future Vision...")
    build_slide_7_team_vision(prs)

    prs.save(output_path)
    print(f"✅ Presentation successfully generated and saved to: {output_path}")

# ==============================================================================
# SLIDE 1: الغلاف (Cover)
# ==============================================================================
def build_slide_1_cover(prs):
    slide = create_full_slide(prs, COLOR_NAVY_DARK)

    # Decorative geometric Islamic corner element
    accent_bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(-1.2), Inches(4.2), Inches(4.2)
    )
    accent_bg.fill.solid()
    accent_bg.fill.fore_color.rgb = RGBColor(22, 34, 58)
    accent_bg.line.fill.background()

    # Event Badge Pill at Top Right
    pill_w = Inches(5.6)
    pill_h = Inches(0.42)
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.933), Inches(0.75), pill_w, pill_h
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(22, 36, 64)
    pill.line.color.rgb = COLOR_EMERALD
    pill.line.width = Pt(1.2)
    tf_pill = pill.text_frame
    tf_pill.margin_left = tf_pill.margin_right = tf_pill.margin_top = tf_pill.margin_bottom = 0
    p_pill = tf_pill.paragraphs[0]
    p_pill.text = "✦  تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي 2026"
    set_center(p_pill)
    p_pill.font.name = FONT_HEADING
    p_pill.font.size = Pt(11)
    p_pill.font.bold = True
    p_pill.font.color.rgb = COLOR_EMERALD_BRIGHT

    # Main Hero Container Box
    hero_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(3.3))
    tf_hero = hero_box.text_frame
    tf_hero.word_wrap = True
    tf_hero.margin_left = tf_hero.margin_right = tf_hero.margin_top = tf_hero.margin_bottom = 0

    # 1. Brand Title (Arabic)
    p_brand = tf_hero.paragraphs[0]
    p_brand.text = "بيّنة AI"
    set_rtl(p_brand)
    p_brand.font.name = FONT_HEADING
    p_brand.font.size = Pt(44)
    p_brand.font.bold = True
    p_brand.font.color.rgb = COLOR_WHITE
    p_brand.space_after = Pt(0)

    # 2. Brand Subtitle (English Engine name)
    p_eng = tf_hero.add_paragraph()
    p_eng.text = "Bayyinah Engine — Grounded Verification System"
    set_rtl(p_eng)
    p_eng.font.name = FONT_LATIN
    p_eng.font.size = Pt(18)
    p_eng.font.bold = True
    p_eng.font.color.rgb = RGBColor(148, 163, 184)
    p_eng.space_after = Pt(12)

    # 3. Main Subtitle (Arabic)
    p_sub = tf_hero.add_paragraph()
    p_sub.text = "محرك التحقق المعرفي وتتبع الأدلة الشرعية لتمكين المعرّفين بالإسلام"
    set_rtl(p_sub)
    p_sub.font.name = FONT_HEADING
    p_sub.font.size = Pt(21)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_EMERALD_BRIGHT
    p_sub.space_after = Pt(12)

    # 4. Track Info
    p_track = tf_hero.add_paragraph()
    p_track.text = "المسار التنافسي: أدوات المعرفة والتحقق لتمكين المعرفين بالإسلام"
    set_rtl(p_track)
    p_track.font.name = FONT_BODY
    p_track.font.size = Pt(14)
    p_track.font.color.rgb = RGBColor(203, 213, 225)
    p_track.space_after = Pt(6)

    # 5. Tagline
    p_tagline = tf_hero.add_paragraph()
    p_tagline.text = "منظومة استدلال مقيد (Grounded RAG) مع إسناد لحظي للأحاديث والمتون المحققة وبروتوكول امتناع صارم ضد الهلوسة"
    set_rtl(p_tagline)
    p_tagline.font.name = FONT_BODY
    p_tagline.font.size = Pt(12.5)
    p_tagline.font.color.rgb = RGBColor(148, 163, 184)

    # 3 Bottom Feature Preview Cards
    features = [
        ("التحقق الفوري والإسناد", "ربط مباشر بين كل ادعاء أو اقتباس ونص الشاهد ورقم الحديث من أمهات المصادر المعتمدة.", "01"),
        ("صفر هلوسة معرفية", "حصر التوليد في سياق المراجع المحققة مع بروتوكول الامتناع الذكي عند عدم التيقن.", "02"),
        ("تمكين ميداني للدعاة", "بطاقة دليل موثقة (Evidence Card) جاهزة للاستخدام اللحظي في الحوارات والمناظرات.", "03"),
    ]

    card_w = Inches(3.68)
    card_h = Inches(1.75)
    gap = Inches(0.34)
    start_y = Inches(4.88)

    for idx, (ft_title, ft_desc, ft_num) in enumerate(features):
        pos_from_right = idx
        card_x = Inches(0.8) + (2 - pos_from_right) * (card_w + gap)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, start_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_NAVY_CARD
        card.line.color.rgb = COLOR_NAVY_BORDER
        card.line.width = Pt(1)

        # Top Emerald indicator
        top_line = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.2), start_y, card_w - Inches(0.4), Inches(0.04)
        )
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = COLOR_EMERALD_BRIGHT
        top_line.line.fill.background()

        tb = slide.shapes.add_textbox(card_x + Inches(0.22), start_y + Inches(0.18), card_w - Inches(0.44), card_h - Inches(0.28))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = f"{ft_title}  •  {ft_num}"
        set_rtl(p1)
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = ft_desc
        set_rtl(p2)
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = RGBColor(148, 163, 184)

    # Slide 1 Footer
    foot_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.35))
    tf_f = foot_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "وثيقة العرض الفني والحل التقني  |  النسخة الموجهة للتحكيم  |  أكتوبر 2026"
    set_center(p_f)
    p_f.font.name = FONT_BODY
    p_f.font.size = Pt(10.5)
    p_f.font.color.rgb = RGBColor(100, 116, 139)

# ==============================================================================
# SLIDE 2: المشكلة وسياق الحاجة (Problem & Context)
# ==============================================================================
def build_slide_2_problem(prs):
    slide = create_full_slide(prs, COLOR_BG_LIGHT)
    add_header(
        slide,
        tag_text="سياق الحاجة  •  المشهد الميداني الحالي",
        title_text="التحديات الواقعية في الميدان المعرفي والدعوي",
        subtitle_text="فجوات حرجة تعيق المعرّفين وصنّاع المحتوى وتفتح الباب لتسرب الروايات غير المحققة"
    )

    cards_data = [
        {
            "num": "01",
            "tag": "تحدي الكفاءة والوقت",
            "title": "صعوبة وبطء التحقق الميداني",
            "subtitle": "Manual Verification Overhead",
            "points": [
                "استنزاف وقت الباحث والداعية في البحث اليدوي الشاق بين أمهات الكتب والموسوعات.",
                "صعوبة التثبت اللحظي من درجة وصحة الأحاديث ورتبة الرواة أثناء الحوارات المباشرة.",
                "تأخر الرد على التساؤلات والشبهات مما يؤدي إلى ضياع فرصة التأثير الإيجابي في المتلقي."
            ],
            "impact": "الأثر: هدر أكثر من 60% من وقت الباحثين في التخريج اليدوي والتحقق المكتبي."
        },
        {
            "num": "02",
            "tag": "خطر السلامة المعرفية",
            "title": "مخاطر الهلوسة الرقمية",
            "subtitle": "AI Hallucination & Plausible Fiction",
            "points": [
                "توليد نماذج اللغة العامة (LLMs) لأحاديث وروايات وتفاسير مختلقة كلياً أو ناقصة.",
                "نسبة أقوال وأحكام لمصادر شرعية غير صحيحة بصياغة لغوية خادعة توهم بالموثوقية.",
                "افتقار النماذج التجارية للمعاجم التخصصية والتمييز الدقيق بين الروايات الصحيحة والضعيفة."
            ],
            "impact": "الأثر: تضليل معرفي خطير عند الاعتماد على الذكاء الاصطناعي العام دون رقابة وضبط."
        },
        {
            "num": "03",
            "tag": "فجوة التوثيق والمصدر",
            "title": "غياب أدوات التتبع اللحظي",
            "subtitle": "Lack of Real-Time Grounding Tools",
            "points": [
                "افتقار المحتوى الرقمي الإسلامي إلى أداة تتبع مباشر تُظهر نص الشاهد والمصدر بدقة.",
                "غياب الربط الآلي برقم الحديث، الباب، وحكم جهابذة المحدثين المعتمدين.",
                "صعوبة التحقق من صحة المقاطع والنصوص المتداولة على شبكات التواصل الاجتماعي فورياً."
            ],
            "impact": "الأثر: انعدام الثقة في المحتوى المنشور وصعوبة إسناد الحجج في النقاشات الفكرية."
        },
    ]

    card_w = Inches(3.68)
    card_h = Inches(4.35)
    gap = Inches(0.34)
    start_y = Inches(2.1)

    for idx, c in enumerate(cards_data):
        pos_from_right = idx
        card_x = Inches(0.8) + (2 - pos_from_right) * (card_w + gap)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, start_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_LIGHT
        card.line.color.rgb = COLOR_BORDER_LIGHT
        card.line.width = Pt(1.2)

        # Top Badge Ribbon
        ribbon = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.2), start_y + Inches(0.18), card_w - Inches(0.4), Inches(0.36)
        )
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = COLOR_AMBER_BG if idx == 1 else COLOR_EMERALD_BG
        ribbon.line.fill.background()
        p_rib = ribbon.text_frame.paragraphs[0]
        p_rib.text = f"{c['num']}  •  {c['tag']}"
        set_center(p_rib)
        p_rib.font.name = FONT_HEADING
        p_rib.font.size = Pt(11)
        p_rib.font.bold = True
        p_rib.font.color.rgb = COLOR_AMBER if idx == 1 else COLOR_EMERALD

        # Text inside card
        tb = slide.shapes.add_textbox(
            card_x + Inches(0.22), start_y + Inches(0.66), card_w - Inches(0.44), card_h - Inches(0.78)
        )
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_title = tf.paragraphs[0]
        p_title.text = c["title"]
        set_rtl(p_title)
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_DARK

        p_sub = tf.add_paragraph()
        p_sub.text = c["subtitle"]
        set_rtl(p_sub)
        p_sub.font.name = FONT_LATIN
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.space_after = Pt(8)

        for pt in c["points"]:
            p_pt = tf.add_paragraph()
            p_pt.text = f"•  {pt}"
            set_rtl(p_pt)
            p_pt.font.name = FONT_BODY
            p_pt.font.size = Pt(11.5)
            p_pt.font.color.rgb = COLOR_TEXT_BODY
            p_pt.space_after = Pt(6)

        # Bottom Impact Box
        impact_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.18), start_y + card_h - Inches(0.78), card_w - Inches(0.36), Inches(0.62)
        )
        impact_box.fill.solid()
        impact_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
        impact_box.line.color.rgb = COLOR_BORDER_LIGHT
        impact_box.line.width = Pt(0.8)
        tf_imp = impact_box.text_frame
        tf_imp.word_wrap = True
        tf_imp.margin_left = tf_imp.margin_right = Inches(0.1)
        tf_imp.margin_top = Inches(0.06)
        p_imp = tf_imp.paragraphs[0]
        p_imp.text = c["impact"]
        set_rtl(p_imp)
        p_imp.font.name = FONT_BODY
        p_imp.font.size = Pt(10)
        p_imp.font.bold = True
        p_imp.font.color.rgb = COLOR_TEXT_DARK

    # Bottom summary banner
    summary_bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.3)
    )
    summary_bg.fill.solid()
    summary_bg.fill.fore_color.rgb = COLOR_EMERALD_BG
    summary_bg.line.fill.background()
    p_sum = summary_bg.text_frame.paragraphs[0]
    p_sum.text = "الخلاصة: الميدان المعرفي بحاجة إلى محرك تحقق هندسي محكم لا يعتمد على التوليد المفتوح، بل يربط كل نص بمصدره المعتمد."
    set_center(p_sum)
    p_sum.font.name = FONT_BODY
    p_sum.font.size = Pt(10.5)
    p_sum.font.bold = True
    p_sum.font.color.rgb = COLOR_EMERALD

    add_footer(slide, 2)

# ==============================================================================
# SLIDE 3: الحل المبتكر والقيمة المضافة (Solution & Value Proposition)
# ==============================================================================
def build_slide_3_solution(prs):
    slide = create_full_slide(prs, COLOR_BG_LIGHT)
    add_header(
        slide,
        tag_text="الحل المبتكر  •  القيمة التنافسية المضافة",
        title_text="الحل: محرك استدلال وتحقق معزز بالاسترجاع المقيد",
        subtitle_text="Grounded RAG — منظومة هندسية ذكية تدمج الضبط الشرعي الصارم مع البحث الدلالي والهجين"
    )

    pillars = [
        {
            "tag": "الركيزة الأولى",
            "title": "التحقق الفوري والإسناد",
            "subtitle": "Instant Verification & Grounding",
            "desc": "استخراج الادعاءات والشواهد من استفسار المستخدم ومطابقتها دلالياً ولفظياً مع المراجع الحقلية المعتمدة فورياً.",
            "features": [
                "تجزئة آلية دقيقة للنصوص وتحديد مواضع الاستدلال الشرعي.",
                "مطابقة هجينة فائقة السرعة مع متون الحديث المعتمدة والمحققة.",
                "كشف التصحيف أو التحريف اللفظي الطفيف في المتن بدقة عالية."
            ],
            "kpi_label": "سرعة الاستجابة",
            "kpi_val": "أقل من 800ms"
        },
        {
            "tag": "الركيزة الثانية",
            "title": "بطاقة الدليل الموثقة",
            "subtitle": "Evidence Card Architecture",
            "desc": "إبراز حالة الدليل ودرجته وسنده ورقم المرجع بشكل شفاف، مرئي، وقابل للتتبع الفوري بنقرة واحدة.",
            "features": [
                "عرض حكم المحدثين (صحيح، حسن، ضعيف) بترميز لوني واضح.",
                "إبراز المصدر، اسم الكتاب، الباب، ورقم الحديث المعتمد دولياً.",
                "عرض نص الشاهد الأصلي مع سياقه الكامل لمنع الاقتطاع المخل."
            ],
            "kpi_label": "شفافية الإسناد",
            "kpi_val": "100% توثيق قطعي"
        },
        {
            "tag": "الركيزة الثالثة",
            "title": "بروتوكول الامتناع الذكي",
            "subtitle": "Abstention & Human Handoff",
            "desc": "توقف النظام تلقائياً عن التوليد عند انخفاض عتبة اليقين، مع إحالة الاستفسارات المشتبهة لمختص بشري.",
            "features": [
                "تطبيق سياسة صارمة: الصمت والامتناع أفضل من توليد معلومة خاطئة.",
                "حظر التوليد الحر تماماً عند غياب السند المباشر في قاعدة المعرفة.",
                "لوحة تدقيق بشري مدمجة لتحكيم المسائل الدقيقة وتطوير النظام."
            ],
            "kpi_label": "الأمان المعرفي",
            "kpi_val": "0% هلوسة رقمية"
        }
    ]

    card_w = Inches(3.68)
    card_h = Inches(4.35)
    gap = Inches(0.34)
    start_y = Inches(2.1)

    for idx, p in enumerate(pillars):
        pos_from_right = idx
        card_x = Inches(0.8) + (2 - pos_from_right) * (card_w + gap)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, start_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_LIGHT
        card.line.color.rgb = COLOR_EMERALD_BORDER if idx == 0 else COLOR_BORDER_LIGHT
        card.line.width = Pt(1.4 if idx == 0 else 1.0)

        # Top Badge
        tag_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.2), start_y + Inches(0.18), card_w - Inches(0.4), Inches(0.36)
        )
        tag_bg.fill.solid()
        tag_bg.fill.fore_color.rgb = COLOR_EMERALD_BG
        tag_bg.line.fill.background()
        p_tag = tag_bg.text_frame.paragraphs[0]
        p_tag.text = f"✦  {p['tag']}"
        set_center(p_tag)
        p_tag.font.name = FONT_HEADING
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_EMERALD

        tb = slide.shapes.add_textbox(card_x + Inches(0.22), start_y + Inches(0.66), card_w - Inches(0.44), card_h - Inches(0.78))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        # Title
        p_title = tf.paragraphs[0]
        p_title.text = p["title"]
        set_rtl(p_title)
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_DARK

        # Latin Subtitle
        p_sub = tf.add_paragraph()
        p_sub.text = p["subtitle"]
        set_rtl(p_sub)
        p_sub.font.name = FONT_LATIN
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.space_after = Pt(8)

        # Desc
        p_desc = tf.add_paragraph()
        p_desc.text = p["desc"]
        set_rtl(p_desc)
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(11.5)
        p_desc.font.color.rgb = COLOR_TEXT_BODY
        p_desc.space_after = Pt(10)

        # Feature items
        for feat in p["features"]:
            p_f = tf.add_paragraph()
            p_f.text = f"✔  {feat}"
            set_rtl(p_f)
            p_f.font.name = FONT_BODY
            p_f.font.size = Pt(11)
            p_f.font.color.rgb = COLOR_TEXT_BODY
            p_f.space_after = Pt(5)

        # Bottom KPI Highlight Box
        kpi_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.18), start_y + card_h - Inches(0.68), card_w - Inches(0.36), Inches(0.52)
        )
        kpi_box.fill.solid()
        kpi_box.fill.fore_color.rgb = COLOR_NAVY_DARK
        kpi_box.line.fill.background()
        tf_k = kpi_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_right = Inches(0.08)
        tf_k.margin_top = Inches(0.06)

        p_k = tf_k.paragraphs[0]
        p_k.text = f"{p['kpi_label']}  •  {p['kpi_val']}"
        set_center(p_k)
        p_k.font.name = FONT_HEADING
        p_k.font.size = Pt(11)
        p_k.font.bold = True
        p_k.font.color.rgb = COLOR_EMERALD_BRIGHT

    add_footer(slide, 3)

# ==============================================================================
# SLIDE 4: المعمارية التقنية للمنظومة (Technical Architecture & Pipeline)
# ==============================================================================
def build_slide_4_architecture(prs):
    slide = create_full_slide(prs, COLOR_BG_LIGHT)
    add_header(
        slide,
        tag_text="المعمارية الهندسية  •  مسار تدفق البيانات",
        title_text="المعمارية الهندسية ومسار البيانات المتكامل",
        subtitle_text="End-to-End Pipeline — خط أنابيب برمجي موجه للإنتاج يجمع بين الاسترجاع الهجين والتحكيم فائق الدقة"
    )

    modules = [
        {
            "step": "01",
            "title": "طبقة البيانات المنسقة",
            "sub": "Curated Knowledge Ingestion",
            "points": [
                "تجميع وتدقيق صحاح الحديث (البخاري، مسلم، والسنن).",
                "كتب التفسير المحققة والمتون الفقهية المعتمدة.",
                "تقسيم النصوص وتوليد التضمينات (Domain Embeddings)."
            ],
            "tech": "PostgreSQL • Normalized DB"
        },
        {
            "step": "02",
            "title": "محرك البحث الهجين",
            "sub": "Hybrid pgvector + BM25",
            "points": [
                "مطابقة دلالية بالأبعاد المتجهة عبر امتداد pgvector.",
                "بحث معجمي لفظي دقيق (BM25) لمطابقة ألفاظ المتون.",
                "دمج النتائج عبر خوارزمية Reciprocal Rank Fusion."
            ],
            "tech": "pgvector • BM25 • RRF"
        },
        {
            "step": "03",
            "title": "طبقة التحكيم وعتبة الثقة",
            "sub": "Confidence & Re-ranking",
            "points": [
                "إعادة ترتيب المخرجات عبر Cross-Encoder تخصصي.",
                "تطبيق عتبة أمان صارمة (Confidence Threshold).",
                "تفعيل بروتوكول الامتناع الذكي عند غياب السند القطعي."
            ],
            "tech": "Cross-Encoder • Guardrails"
        },
        {
            "step": "04",
            "title": "واجهة التشغيل والـ API",
            "sub": "FastAPI & Interactive UI",
            "points": [
                "خدمات برمجية سريعة وموثقة بالكامل عبر FastAPI.",
                "واجهة استعراض تفاعلية لبطاقة الدليل (Evidence Card).",
                "تصدير فوري للاستشهادات والتخريج الشرعي."
            ],
            "tech": "FastAPI • Web UI • REST API"
        }
    ]

    card_w = Inches(2.74)
    card_h = Inches(4.35)
    gap = Inches(0.25)
    start_y = Inches(2.1)

    for idx, m in enumerate(modules):
        pos_from_right = idx
        card_x = Inches(0.8) + (3 - pos_from_right) * (card_w + gap)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, start_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_LIGHT
        card.line.color.rgb = COLOR_BORDER_LIGHT
        card.line.width = Pt(1.1)

        # Top step badge
        step_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.18), start_y + Inches(0.18), card_w - Inches(0.36), Inches(0.4)
        )
        step_bg.fill.solid()
        step_bg.fill.fore_color.rgb = COLOR_NAVY_DARK
        step_bg.line.fill.background()
        p_step = step_bg.text_frame.paragraphs[0]
        p_step.text = f"المرحلة {m['step']}"
        set_center(p_step)
        p_step.font.name = FONT_HEADING
        p_step.font.size = Pt(11.5)
        p_step.font.bold = True
        p_step.font.color.rgb = COLOR_EMERALD_BRIGHT

        tb = slide.shapes.add_textbox(card_x + Inches(0.18), start_y + Inches(0.72), card_w - Inches(0.36), card_h - Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        # Title
        p_title = tf.paragraphs[0]
        p_title.text = m["title"]
        set_rtl(p_title)
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_DARK
        p_title.space_after = Pt(2)

        # Subtitle
        p_sub = tf.add_paragraph()
        p_sub.text = m["sub"]
        set_rtl(p_sub)
        p_sub.font.name = FONT_LATIN
        p_sub.font.size = Pt(9)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.space_after = Pt(10)

        # Points
        for pt in m["points"]:
            p_pt = tf.add_paragraph()
            p_pt.text = f"•  {pt}"
            set_rtl(p_pt)
            p_pt.font.name = FONT_BODY
            p_pt.font.size = Pt(10.8)
            p_pt.font.color.rgb = COLOR_TEXT_BODY
            p_pt.space_after = Pt(6)

        # Bottom Tech Tag
        tech_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.15), start_y + card_h - Inches(0.65), card_w - Inches(0.3), Inches(0.48)
        )
        tech_box.fill.solid()
        tech_box.fill.fore_color.rgb = COLOR_BLUE_BG
        tech_box.line.color.rgb = COLOR_BLUE_BORDER
        tech_box.line.width = Pt(0.8)
        p_t = tech_box.text_frame.paragraphs[0]
        p_t.text = m["tech"]
        set_center(p_t)
        p_t.font.name = FONT_LATIN
        p_t.font.size = Pt(9.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_BLUE_TEXT

    add_footer(slide, 4)

# ==============================================================================
# SLIDE 5: الموثوقية والسلامة العلمية (Reliability & Scientific Safety)
# ==============================================================================
def build_slide_5_safety(prs):
    slide = create_full_slide(prs, COLOR_BG_LIGHT)
    add_header(
        slide,
        tag_text="الموثوقية والضوابط  •  معايير الجودة الشرعية",
        title_text="معايير السلامة المعرفية ومكافحة الخطأ",
        subtitle_text="منظومة حوكمة هندسية متشددة لضمان أعلى مستويات الدقة وحماية المحتوى الإسلامي من التضليل"
    )

    # 3 Top KPI Cards
    kpis = [
        ("نسبة الهلوسة الرقمية", "0%", "تصفير التوليد الحر عبر تقييد النموذج بالنصوص المعتمدة حصراً"),
        ("دقة استرجاع المصدر", "+95%", "مطابقة دلالية وسياقية فائقة لمتون الأحاديث والروايات المعتمدة"),
        ("شفافية الإسناد والعزو", "100%", "كل مخرج مدعوم برقم الحديث، اسم الكتاب، وتخريج المحدثين"),
    ]

    kpi_w = Inches(3.68)
    kpi_h = Inches(1.35)
    gap = Inches(0.34)
    start_y_kpi = Inches(2.05)

    for idx, (title, num, desc) in enumerate(kpis):
        pos_from_right = idx
        kpi_x = Inches(0.8) + (2 - pos_from_right) * (kpi_w + gap)

        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, kpi_x, start_y_kpi, kpi_w, kpi_h)
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_DARK
        box.line.color.rgb = COLOR_EMERALD
        box.line.width = Pt(1.2)

        tb = slide.shapes.add_textbox(kpi_x + Inches(0.18), start_y_kpi + Inches(0.1), kpi_w - Inches(0.36), kpi_h - Inches(0.16))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        # Tier 1: Stat Number
        p_num = tf.paragraphs[0]
        p_num.text = num
        set_rtl(p_num)
        p_num.font.name = FONT_LATIN
        p_num.font.size = Pt(22)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_EMERALD_BRIGHT

        # Tier 2: Metric Title
        p_title = tf.add_paragraph()
        p_title.text = title
        set_rtl(p_title)
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(12)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE
        p_title.space_after = Pt(2)

        # Tier 3: Description
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        set_rtl(p_desc)
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = RGBColor(203, 213, 225)

    # 2 Detailed Strategy Cards below
    strat_w = Inches(5.69)
    strat_h = Inches(3.1)
    gap_strat = Inches(0.35)
    start_y_strat = Inches(3.6)

    strategies = [
        {
            "tag": "المعيار الأول • الحظر المطلق للتوليد الحر",
            "title": "حصر المخرجات في سياق المراجع المعتمدة",
            "sub": "Zero Unbounded Generation Policy",
            "bullets": [
                "منع نماذج الذكاء الاصطناعي من التخمين أو التوليد غير المستند إلى وثائق مسترجعة.",
                "نظام فحص الحقائق المزدوج: مقارنة كل جملة مصدرة بسياق المتن المحقق حرفياً.",
                "في حال عدم وجود تطابق بنسبة تتجاوز عتبة الأمان، يتم تفعيل بروتوكول الامتناع فوراً.",
                "استبعاد الآراء الشاذة والروايات الضعيفة والموضوعة من نطاق الاستدلال المعتمد."
            ]
        },
        {
            "tag": "المعيار الثاني • الحوكمة والتحكيم التخصصي",
            "title": "مسار المراجعة والتدقيق البشري المدمج",
            "sub": "Human-in-the-Loop Verification Pipeline",
            "bullets": [
                "لوحة مراجعة للمختصين الشرعيين لتقييم الحالات التي تقع في منطقة الشك المعرفي.",
                "حلقة تغذية راجعة مستمرة (Feedback Loop) لإعادة معايرة أوزان التضمينات وتحسين الدقة.",
                "توثيق سجل تدقيق غير قابل للتعديل (Audit Trail) يوضح كل استعلام ومصدره وحكمه.",
                "تمكين المؤسسات الدعوية من تخصيص معايير القبول والترجيح حسب سياساتها المعتمدة."
            ]
        }
    ]

    for idx, st in enumerate(strategies):
        pos_from_right = idx
        st_x = Inches(0.8) + (1 - pos_from_right) * (strat_w + gap_strat)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, st_x, start_y_strat, strat_w, strat_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_LIGHT
        card.line.color.rgb = COLOR_BORDER_LIGHT
        card.line.width = Pt(1.1)

        # Header tag
        tag_b = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, st_x + Inches(0.2), start_y_strat + Inches(0.18), strat_w - Inches(0.4), Inches(0.32)
        )
        tag_b.fill.solid()
        tag_b.fill.fore_color.rgb = COLOR_EMERALD_BG
        tag_b.line.fill.background()
        p_tb = tag_b.text_frame.paragraphs[0]
        p_tb.text = st["tag"]
        set_center(p_tb)
        p_tb.font.name = FONT_HEADING
        p_tb.font.size = Pt(10.5)
        p_tb.font.bold = True
        p_tb.font.color.rgb = COLOR_EMERALD

        tb = slide.shapes.add_textbox(st_x + Inches(0.25), start_y_strat + Inches(0.58), strat_w - Inches(0.5), strat_h - Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = st["title"]
        set_rtl(p_t)
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(14.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_DARK

        p_s = tf.add_paragraph()
        p_s.text = st["sub"]
        set_rtl(p_s)
        p_s.font.name = FONT_LATIN
        p_s.font.size = Pt(9.5)
        p_s.font.color.rgb = COLOR_TEXT_MUTED
        p_s.space_after = Pt(8)

        for b in st["bullets"]:
            p_b = tf.add_paragraph()
            p_b.text = f"✔  {b}"
            set_rtl(p_b)
            p_b.font.name = FONT_BODY
            p_b.font.size = Pt(11.2)
            p_b.font.color.rgb = COLOR_TEXT_BODY
            p_b.space_after = Pt(5)

    add_footer(slide, 5)

# ==============================================================================
# SLIDE 6: خطة التنفيذ خلال أيام التحدي 4–6 أكتوبر (Hackathon Roadmap)
# ==============================================================================
def build_slide_6_roadmap(prs):
    slide = create_full_slide(prs, COLOR_BG_LIGHT)
    add_header(
        slide,
        tag_text="خارطة الطريق  •  مراحل الإنجاز السريع",
        title_text="خطة التنفيذ لبناء المنتج الأولي الجاهز",
        subtitle_text="Production-Ready MVP Roadmap — جدول زمني مكثف ومحكم لتحقيق أعلى مخرجات هندسية (4–6 أكتوبر)"
    )

    days = [
        {
            "day": "اليوم الأول (4 أكتوبر)",
            "theme": "التأسيس والبيانات والبنية الخلفية",
            "tag": "Day 1 • Ingestion & Infrastructure",
            "tasks": [
                "تجهيز وهيكلة قاعدة بيانات المتون المحققة والتفاسير المعتمدة.",
                "توليد التضمينات الدلالية وحفظ الفهارس المتجهة عبر pgvector.",
                "بناء مسارات الـ API الأساسية وخدمات الاستعلام عبر FastAPI.",
                "اختبار وتأكيد زمن الاستجابة الأولي للاسترجاع المتجه."
            ],
            "deliverable": "المخرج: قاعدة معرفية مهيكلة وخادم API يعمل بنجاح."
        },
        {
            "day": "اليوم الثاني (5 أكتوبر)",
            "theme": "محرك البحث والامتناع والسلامة",
            "tag": "Day 2 • Hybrid Search & Guardrails",
            "tasks": [
                "دمج البحث اللفظي BM25 مع البحث الدلالي عبر خوارزمية RRF.",
                "تفعيل طبقة إعادة الترتيب (Re-ranking) وتحديد عتبة الثقة.",
                "برمجة منطق بروتوكول الامتناع الذكي وإحالة الحالات المشتبهة.",
                "إجراء اختبارات دقة الاسترجاع والسلامة المعرفية ومكافحة الهلوسة."
            ],
            "deliverable": "المخرج: محرك استدلال دقيق مع بروتوكول أمان معتمد."
        },
        {
            "day": "اليوم الثالث (6 أكتوبر)",
            "theme": "الواجهة التفاعلية والتوثيق والإطلاق",
            "tag": "Day 3 • Live Demo & Submission",
            "tasks": [
                "إطلاق الواجهة التفاعلية (Interactive UI) وبطاقة الدليل الموثقة.",
                "توثيق المستودع البرمجي على GitHub مع تعليمات التشغيل السريع.",
                "تسجيل الفيديو التوضيحي للحل واستعراض حالات الاستخدام الحية.",
                "مراجعة معايير التحكيم وضمان الجاهزية التامة للعرض النهائي."
            ],
            "deliverable": "المخرج: نموذج حي كامل جاهز لاختبارات لجنة التحكيم."
        },
    ]

    card_w = Inches(3.68)
    card_h = Inches(4.35)
    gap = Inches(0.34)
    start_y = Inches(2.1)

    for idx, d in enumerate(days):
        pos_from_right = idx
        card_x = Inches(0.8) + (2 - pos_from_right) * (card_w + gap)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, start_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_LIGHT
        card.line.color.rgb = COLOR_BORDER_LIGHT
        card.line.width = Pt(1.2)

        # Header Ribbon
        ribbon = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.18), start_y + Inches(0.18), card_w - Inches(0.36), Inches(0.4)
        )
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = COLOR_NAVY_DARK
        ribbon.line.fill.background()
        p_r = ribbon.text_frame.paragraphs[0]
        p_r.text = d["day"]
        set_center(p_r)
        p_r.font.name = FONT_HEADING
        p_r.font.size = Pt(12)
        p_r.font.bold = True
        p_r.font.color.rgb = COLOR_EMERALD_BRIGHT

        tb = slide.shapes.add_textbox(card_x + Inches(0.22), start_y + Inches(0.68), card_w - Inches(0.44), card_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_th = tf.paragraphs[0]
        p_th.text = d["theme"]
        set_rtl(p_th)
        p_th.font.name = FONT_HEADING
        p_th.font.size = Pt(14)
        p_th.font.bold = True
        p_th.font.color.rgb = COLOR_TEXT_DARK
        p_th.space_after = Pt(2)

        p_tg = tf.add_paragraph()
        p_tg.text = d["tag"]
        set_rtl(p_tg)
        p_tg.font.name = FONT_LATIN
        p_tg.font.size = Pt(9.5)
        p_tg.font.color.rgb = COLOR_TEXT_MUTED
        p_tg.space_after = Pt(10)

        for t in d["tasks"]:
            p_t = tf.add_paragraph()
            p_t.text = f"✔  {t}"
            set_rtl(p_t)
            p_t.font.name = FONT_BODY
            p_t.font.size = Pt(11)
            p_t.font.color.rgb = COLOR_TEXT_BODY
            p_t.space_after = Pt(6)

        # Deliverable Box
        deliv_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.18), start_y + card_h - Inches(0.72), card_w - Inches(0.36), Inches(0.55)
        )
        deliv_box.fill.solid()
        deliv_box.fill.fore_color.rgb = COLOR_EMERALD_BG
        deliv_box.line.color.rgb = COLOR_EMERALD_BORDER
        deliv_box.line.width = Pt(0.8)
        tf_del = deliv_box.text_frame
        tf_del.word_wrap = True
        tf_del.margin_left = tf_del.margin_right = Inches(0.1)
        tf_del.margin_top = Inches(0.06)
        p_d = tf_del.paragraphs[0]
        p_d.text = d["deliverable"]
        set_rtl(p_d)
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(10.5)
        p_d.font.bold = True
        p_d.font.color.rgb = COLOR_EMERALD

    add_footer(slide, 6)

# ==============================================================================
# SLIDE 7: الفريق والرؤية المستقبلية (Team & Future Vision)
# ==============================================================================
def build_slide_7_team_vision(prs):
    slide = create_full_slide(prs, COLOR_NAVY_DARK)
    add_header(
        slide,
        tag_text="القدرة التنفيذية  •  الاستدامة والتوسع",
        title_text="القدرة التنفيذية والتوسع المستدام",
        subtitle_text="فريق متخصص بتكامل هندسي وعلمي، ورؤية طموحة لتحويل الفكرة إلى بنية تحتية دائمة",
        is_dark=True
    )

    cols = [
        {
            "tag": "الكفاءة التنفيذية",
            "title": "الفريق وتكامل الخبرات",
            "sub": "Execution Capability & Readiness",
            "points": [
                "هندسة البرمجيات والأنظمة الموزعة: تصميم واجهات خلفية عالية الأداء وقابلة للتوسع اللامحدود.",
                "هندسة الذكاء الاصطناعي وRAG: تخصص دقيق في معالجة اللغات الطبيعية والتضمينات المتجهة.",
                "إدارة قواعد البيانات الحديثة: خبرة عملية في إدارة PostgreSQL وpgvector وفهارس البحث.",
                "الاستيعاب الشرعي والمنهجي: حرص صارم على مراعاة أصول التخريج والتحقيق العلمي المعتمد."
            ]
        },
        {
            "tag": "الاستدامة بعد التحدي",
            "title": "التحقق كخدمة برمجية",
            "sub": "Verification-as-a-Service (VaaS)",
            "points": [
                "توفير المحرك كخدمة برمجية سحابية (API) قابلة للربط السهل مع المنصات والمواقع الإسلامية.",
                "إتاحة إضافة للمتصفح (Chrome/Edge Extension) تتيح التحقق الفوري من أي محتوى بنقرة زر.",
                "درع أمان ذكي (Guardrail Plugin) لروبوتات المحادثة الإسلامية لمنع توليد الروايات الخاطئة.",
                "نموذج تشغيلي مستدام عبر شراكات استراتيجية مع المراكز الدعوية والمؤسسات الرقمية."
            ]
        },
        {
            "tag": "آفاق التوسع المعرفي",
            "title": "الرؤية المستقبلية للمنظومة",
            "sub": "Future Roadmap & Scaling",
            "points": [
                "توسيع قاعدة المعرفة لتشمل فتاوى المجامع الفقهية المعتمدة والمعاجم اللغوية الموسعة.",
                "دعم اللغات العالمية الحية (الإنجليزية، الفرنسية، الأردية، الإندونيسية) لخدمة دعاة المهجر.",
                "تضمين المقارنة النصية بين طبعات وروايات المتون المعتمدة لتوفير أعلى مستويات التوثيق.",
                "بناء مرصد رقمي لرصد ومكافحة الأحاديث الموضوعة والشبهات المتداولة لحظياً."
            ]
        }
    ]

    card_w = Inches(3.68)
    card_h = Inches(4.35)
    gap = Inches(0.34)
    start_y = Inches(2.1)

    for idx, c in enumerate(cols):
        pos_from_right = idx
        card_x = Inches(0.8) + (2 - pos_from_right) * (card_w + gap)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, start_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_NAVY_CARD
        card.line.color.rgb = COLOR_NAVY_BORDER
        card.line.width = Pt(1.1)

        # Top Accent Ribbon
        ribbon = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, card_x + Inches(0.18), start_y + Inches(0.18), card_w - Inches(0.36), Inches(0.36)
        )
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = RGBColor(22, 36, 64)
        ribbon.line.color.rgb = COLOR_EMERALD_BRIGHT
        ribbon.line.width = Pt(1)
        p_r = ribbon.text_frame.paragraphs[0]
        p_r.text = f"✦  {c['tag']}"
        set_center(p_r)
        p_r.font.name = FONT_HEADING
        p_r.font.size = Pt(11)
        p_r.font.bold = True
        p_r.font.color.rgb = COLOR_EMERALD_BRIGHT

        tb = slide.shapes.add_textbox(card_x + Inches(0.22), start_y + Inches(0.68), card_w - Inches(0.44), card_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_title = tf.paragraphs[0]
        p_title.text = c["title"]
        set_rtl(p_title)
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(15.5)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE
        p_title.space_after = Pt(2)

        p_sub = tf.add_paragraph()
        p_sub.text = c["sub"]
        set_rtl(p_sub)
        p_sub.font.name = FONT_LATIN
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = RGBColor(148, 163, 184)
        p_sub.space_after = Pt(10)

        for pt in c["points"]:
            p_pt = tf.add_paragraph()
            p_pt.text = f"•  {pt}"
            set_rtl(p_pt)
            p_pt.font.name = FONT_BODY
            p_pt.font.size = Pt(11)
            p_pt.font.color.rgb = RGBColor(226, 232, 240)
            p_pt.space_after = Pt(6)

    # Bottom Closing Slogan
    slogan = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.52), Inches(11.733), Inches(0.32)
    )
    slogan.fill.solid()
    slogan.fill.fore_color.rgb = RGBColor(22, 36, 64)
    slogan.line.color.rgb = COLOR_EMERALD
    slogan.line.width = Pt(0.8)
    p_slog = slogan.text_frame.paragraphs[0]
    p_slog.text = "بيّنة AI: نحو محتوى إسلامي موثق، رصين، ومحصّن ضد الهلوسة بأحدث حلول الذكاء الاصطناعي."
    set_center(p_slog)
    p_slog.font.name = FONT_HEADING
    p_slog.font.size = Pt(11)
    p_slog.font.bold = True
    p_slog.font.color.rgb = COLOR_EMERALD_BRIGHT

    add_footer(slide, 7, is_dark=True)

# ==============================================================================
# MAIN RUNNER
# ==============================================================================
if __name__ == "__main__":
    output_filename = "Bayyinah_AI_Presentation.pptx"
    if len(sys.argv) > 1:
        output_filename = sys.argv[1]
    build_presentation(output_filename)
